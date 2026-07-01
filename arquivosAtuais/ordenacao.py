import time
from openpyxl import load_workbook, Workbook
from selenium import webdriver
from selenium.webdriver.common.by import By

def tratar_nome_autor(autor_string):
    if not autor_string:
        return ""
    for separador in [";", ",", " e "]:
        if separador in autor_string:
            autor_string = autor_string.split(separador)[0].strip()
            break
    partes = autor_string.strip().split()
    if len(partes) <= 1:
        return autor_string
    sobrenome = partes[-1]       
    nome = " ".join(partes[:-1]) 
    return f"{sobrenome}, {nome}"

# 1. CARREGAR PLANILHA ORIGINAL
print("[INFO] Lendo a planilha original...")
# wb = load_workbook("planilhaFinal.xlsx")
wb = load_workbook("planilha_mais_recente.xlsx")

ws = wb["Sheet1"]

# 2. CONFIGURAR DRIVER DO SELENIUM
options = webdriver.ChromeOptions()
driver = webdriver.Chrome(options=options)

lista_livros = []

print("[INFO] Iniciando a coleta de dados e geração de códigos Cutter...")

for linha in range(3, ws.max_row + 1):
    titulo = ws[f'B{linha}'].value
    if not titulo:
        continue
        
    autor_raw = ws[f'D{linha}'].value or ""
    cdd = ws[f'K{linha}'].value or ""
    
    # Captura os dados originais da linha
    dados_da_linha = [ws.cell(row=linha, column=col).value for col in range(1, ws.max_column + 1)]
    
    codigo_cutter = ""
    autor_formatado = tratar_nome_autor(autor_raw)
    
    if autor_formatado:
        try:
            driver.get("https://www.tabelacutter.com/")
            time.sleep(1)
            
            input_cutter = driver.find_element(By.NAME, "e")
            input_cutter.clear()
            input_cutter.send_keys(autor_formatado)
            
            btn_gerar = driver.find_element(By.XPATH, "//button[contains(text(), 'Gerar código Cutter')]")
            btn_gerar.click()
            time.sleep(1)
            
            elemento_span = driver.find_element(By.ID, "cutterSpan")
            codigo_cutter = elemento_span.text.strip()
            print(f"[LINHA {linha}] {titulo[:30]}... | CDD: {cdd} | Cutter: {codigo_cutter}")
        except Exception as e:
            print(f"[AVISO] Falha ao gerar Cutter na linha {linha}: {e}")
            
    lista_livros.append({
        'cdd': str(cdd).strip() if cdd else "999",
        'cutter': codigo_cutter,
        'dados_originais': dados_da_linha
    })

driver.quit()

# 3. ORDENAÇÃO POR CDD E CUTTER
print("\n[INFO] Ordenando os livros por CDD e Cutter...")
livros_ordenados = sorted(lista_livros, key=lambda x: (x['cdd'], x['cutter']))

# 4. SALVAR OS DADOS ORDENADOS E O CUTTER NA NOVA PLANILHA
print("[INFO] Criando o novo arquivo ordenado...")
wb_novo = Workbook()
ws_novo = wb_novo.active
ws_novo.title = "Sheet1"

# Copia os cabeçalhos originais
for r in range(1, 3):
    for c in range(1, ws.max_column + 1):
        ws_novo.cell(row=r, column=c).value = ws.cell(row=r, column=c).value

# MODIFICAÇÃO: Cria o cabeçalho "Cutter" na coluna L (coluna 12)
coluna_cutter = 12 
ws_novo.cell(row=2, column=coluna_cutter).value = "Cutter"

# Escreve os livros ordenados e insere o Cutter gerado
for idx, livro in enumerate(livros_ordenados, start=3):
    # Escreve os dados que já vinham da planilha antiga
    for col_idx, valor in enumerate(livro['dados_originais'], start=1):
        ws_novo.cell(row=idx, column=col_idx).value = valor
    
    # MODIFICAÇÃO: Grava o Cutter gerado diretamente na nova coluna
    ws_novo.cell(row=idx, column=coluna_cutter).value = livro['cutter']

# Salva o arquivo finalizado
wb_novo.save("planilha_ordenada.xlsx")
print("\n[SUCESSO] Arquivo 'planilha_ordenada.xlsx' gerado com sucesso (Ordenado + Cutter incluído)!")