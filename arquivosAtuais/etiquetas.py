from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, PageBreak
from reportlab.lib import colors

def gerar_pdf_etiquetas():
    print("[INFO] Lendo dados da planilha ordenada...")
    wb = load_workbook("planilha_ordenada.xlsx")
    ws = wb["Sheet1"]

    # 1. COLETAR OS DADOS DA PLANILHA (A partir da linha 3)
    lista_etiquetas = []
    for linha in range(3, ws.max_row + 1):
        titulo = ws.cell(row=linha, column=2).value
        if not titulo:
            continue
            
        subtitulo = ws.cell(row=linha, column=3).value or ""
        cdd = ws.cell(row=linha, column=11).value or ""
        cutter = ws.cell(row=linha, column=12).value or ""
        
        # Junta título e subtítulo de forma elegante
        titulo_completo = f"{titulo} - {subtitulo}" if subtitulo else titulo
        
        lista_etiquetas.append({
            'titulo': titulo_completo,
            'cdd': str(cdd),
            'cutter': str(cutter)
        })

    if not lista_etiquetas:
        print("[AVISO] Nenhuma etiqueta encontrada para gerar.")
        return

    # 2. CONFIGURAR ESTILOS DE TEXTO DO PDF
    styles = getSampleStyleSheet()
    
    # Estilo para o Título (Texto menor para caber no quadrado)
    style_titulo = ParagraphStyle(
        'EstiloTitulo',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=12,
        alignment=TA_CENTER,
        spaceAfter=10
    )
    
    # Estilo para CDD e Cutter (Fontes maiores e destacadas)
    style_codigo = ParagraphStyle(
        'EstiloCodigo',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=16,
        alignment=TA_CENTER
    )

    # 3. CONFIGURAR LAYOUT DA GRADE (3 colunas x 5 linhas por página)
    N_COLUNAS = 3
    N_LINHAS = 5
    ETIQUETAS_POR_PAGINA = N_COLUNAS * N_LINHAS
    
    # Dimensões dos quadrados (em pontos: 1 cm ~ 28.35 pontos)
    largura_coluna = 180  # ~ 6.3 cm
    altura_linha = 145    # ~ 5.1 cm
    
    col_widths = [largura_coluna] * N_COLUNAS
    row_heights = [altura_linha] * N_LINHAS

    # Configuração do documento PDF (A4 com margens pequenas para maximizar espaço)
    doc = SimpleDocTemplate(
        "etiquetas_livros.pdf",
        pagesize=A4,
        leftMargin=25,
        rightMargin=25,
        topMargin=35,
        bottomMargin=25
    )
    
    story = []

    # Agrupa as etiquetas de 15 em 15 (por página)
    for i in range(0, len(lista_etiquetas), ETIQUETAS_POR_PAGINA):
        grupo_pagina = lista_etiquetas[i:i + ETIQUETAS_POR_PAGINA]
        
        # Cria uma matriz vazia 5x3 para a tabela desta página
        matriz_pagina = [[[] for _ in range(N_COLUNAS)] for _ in range(N_LINHAS)]
        
        # Preenche a matriz com o conteúdo das etiquetas
        for idx, item in enumerate(grupo_pagina):
            r = idx // N_COLUNAS
            c = idx % N_COLUNAS
            
            # Monta o conteúdo interno do quadrado
            conteudo_quadrado = [
                Paragraph(item['titulo'], style_titulo),
                Paragraph(item['cdd'], style_codigo),
                Paragraph(item['cutter'], style_codigo)
            ]
            matriz_pagina[r][c] = conteudo_quadrado
            
        # Cria a tabela da página atual
        tabela = Table(matriz_pagina, colWidths=col_widths, rowHeights=row_heights)
        
        # Aplica as bordas dos quadrados e o alinhamento interno vertical
        tabela.setStyle(TableStyle([
            ('GRID', (0,0), (-1,-1), 1, colors.black),      # Cria o contorno dos quadrados
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),          # Centraliza o bloco de texto verticalmente
            ('TOPPADDING', (0,0), (-1,-1), 10),            # Margem interna superior
            ('BOTTOMPADDING', (0,0), (-1,-1), 10),         # Margem interna inferior
            ('LEFTPADDING', (0,0), (-1,-1), 8),            # Margem interna esquerda
            ('RIGHTPADDING', (0,0), (-1,-1), 8),           # Margem interna direita
        ]))
        
        story.append(tabela)
        
        # Se ainda houver mais etiquetas, adiciona uma quebra de página
        if i + ETIQUETAS_POR_PAGINA < len(lista_etiquetas):
            story.append(PageBreak())

    # 4. CONSTRÓI O PDF FINAL
    print("[INFO] Renderizando arquivo PDF...")
    doc.build(story)
    print("[SUCESSO] O arquivo 'etiquetas_livros.pdf' foi gerado e está pronto para impressão!")

if __name__ == "__main__":
    gerar_pdf_etiquetas()