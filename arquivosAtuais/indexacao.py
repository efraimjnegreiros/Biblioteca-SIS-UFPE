# # # # from selenium import webdriver
# # # # from selenium.webdriver.common.by import By
# # # # from selenium.webdriver.common.keys import Keys
# # # # from openpyxl import load_workbook
# # # # import time

# # # # def tratar_nome_autor(autor_string):
# # # #     """
# # # #     Pega o nome do autor da planilha. Se houver mais de um (separado por vírgula, 
# # # #     ponto e vírgula ou 'e'), mantém apenas o primeiro.
# # # #     Inverte o nome para o formato 'Sobrenome, Nome' exigido pelo gerador Cutter.
# # # #     """
# # # #     if not autor_string:
# # # #         return ""
    
# # # #     # Isola o primeiro autor se houver múltiplos (separadores comuns)
# # # #     for separador in [";", ",", " e "]:
# # # #         if separador in autor_string:
# # # #             autor_string = autor_string.split(separador)[0].strip()
# # # #             break
            
# # # #     partes = autor_string.strip().split()
# # # #     if len(partes) <= 1:
# # # #         return autor_string # Retorna o próprio termo se for palavra única
    
# # # #     sobrenome = partes[-1]       # Último nome
# # # #     nome = " ".join(partes[:-1]) # Nomes anteriores
# # # #     return f"{sobrenome}, {nome}"

# # # # # PLANILHA
# # # # wb = load_workbook("planilha.xlsx")
# # # # ws = wb["Sheet1"]

# # # # # CHROME
# # # # options = webdriver.ChromeOptions()
# # # # driver = webdriver.Chrome(options=options)
# # # # driver.maximize_window()

# # # # print("[INFO] Abrindo Biblivre...")
# # # # driver.get("http://localhost/Biblivre5/")
# # # # time.sleep(3)

# # # # # ---------------------------------------------------------------------------
# # # # # PASSO 1: ROTINA DE LOGIN NO BIBLIVRE
# # # # # ---------------------------------------------------------------------------
# # # # try:
# # # #     print("[INFO] Realizando login...")
# # # #     campo_usuario = driver.find_element(By.NAME, "username")
# # # #     campo_usuario.clear()
# # # #     campo_usuario.send_keys("admin")
    
# # # #     campo_senha = driver.find_element(By.NAME, "password")
# # # #     campo_senha.clear()
# # # #     campo_senha.send_keys("Bibsis2025#")
    
# # # #     botao_entrar = driver.find_element(By.XPATH, "//button[contains(@onclick, \"submitForm('login'\")]")
# # # #     botao_entrar.click()
    
# # # #     time.sleep(3)
# # # #     print("[INFO] Login realizado com sucesso!")
# # # # except Exception as e:
# # # #     print(f"[ERRO LOGIN] Não foi possível logar: {e}")
# # # #     driver.save_screenshot("erro_login.png")
# # # #     driver.quit()
# # # #     exit()

# # # # # Guardar o identificador da aba do Biblivre
# # # # aba_biblivre = driver.current_window_handle

# # # # # ---------------------------------------------------------------------------
# # # # # PASSO 2: LAÇO DE REPETIÇÃO DA PLANILHA
# # # # # ---------------------------------------------------------------------------
# # # # # for linha in range(2, 3):
# # # # for linha in range(3, 20):


# # # #     print("\n" + "=" * 60)
# # # #     print(f"[LINHA {linha}] Iniciando processamento")

# # # #     titulo = ws[f'B{linha}'].value or ""
# # # #     subtitulo = ws[f'C{linha}'].value or ""
# # # #     autor_raw = ws[f'D{linha}'].value or ""
# # # #     cdd = ws[f'K{linha}'].value or ""

# # # #     if not titulo:
# # # #         print(f"[LINHA {linha}] Ignorada: Título vazio.")
# # # #         continue

# # # #     print(f"[LINHA {linha}] Livro: {titulo} | Autor original: {autor_raw}")

# # # #     # -----------------------------------------------------------------------
# # # #     # PASSO 2.1: GERAR CÓDIGO CUTTER (Pegando direto do span)
# # # #     # -----------------------------------------------------------------------
# # # #     codigo_autor = ""
# # # #     autor_formatado = tratar_nome_autor(autor_raw)

# # # #     if autor_formatado:
# # # #         try:
# # # #             print(f"[LINHA {linha}] Consultando código Cutter para: '{autor_formatado}'")
# # # #             # Abre uma nova aba para o site do Cutter
# # # #             driver.execute_script("window.open('https://www.tabelacutter.com/', '_blank');")
# # # #             time.sleep(2)
            
# # # #             # Alterna para a aba do Cutter
# # # #             abas = driver.window_handles
# # # #             driver.switch_to.window(abas[-1])
            
# # # #             # Digita no campo de texto do autor
# # # #             input_cutter = driver.find_element(By.NAME, "e")
# # # #             input_cutter.clear()
# # # #             input_cutter.send_keys(autor_formatado)
            
# # # #             # Clica no botão para Gerar o código
# # # #             btn_gerar = driver.find_element(By.XPATH, "//button[contains(text(), 'Gerar código Cutter')]")
# # # #             btn_gerar.click()
# # # #             time.sleep(1.5)
            
# # # #             # MODIFICADO: Captura o texto direto de dentro do <span> com o ID cutterSpan (Ex: N385)
# # # #             elemento_span = driver.find_element(By.ID, "cutterSpan")
# # # #             codigo_autor = elemento_span.text.strip()
            
# # # #             print(f"[LINHA {linha}] Código Cutter obtido direto do span: {codigo_autor}")
            
# # # #             # Fecha a aba do Cutter e retorna para a aba principal do Biblivre
# # # #             driver.close()
# # # #             driver.switch_to.window(aba_biblivre)
# # # #             time.sleep(1)
            
# # # #         except Exception as e_cutter:
# # # #             print(f"[AVISO CUTTER] Falha ao gerar código Cutter na linha {linha}: {e_cutter}")
# # # #             # Em caso de erro, garante que voltou para a aba do Biblivre para não quebrar o fluxo
# # # #             try:
# # # #                 driver.switch_to.window(aba_biblivre)
# # # #             except:
# # # #                 pass

# # # #     # -----------------------------------------------------------------------
# # # #     # PASSO 2.2: PESQUISAR E EDITAR O REGISTRO NO BIBLIVRE
# # # #     # -----------------------------------------------------------------------
# # # #     try:
# # # #         print(f"[LINHA {linha}] Navegando até Catalogação -> Bibliográfica...")
# # # #         menu_catalogacao = driver.find_element(By.CSS_SELECTOR, "li.menu_cataloging")
# # # #         menu_catalogacao.click()
# # # #         time.sleep(0.5)
        
# # # #         submenu_bibliografica = driver.find_element(By.XPATH, "//li[@data-action='cataloging_bibliographic']/a")
# # # #         submenu_bibliografica.click()
# # # #         time.sleep(2.5)

# # # #         print(f"[LINHA {linha}] Pesquisando pelo título: {titulo}")
# # # #         campo_pesquisa = driver.find_element(By.NAME, "query")
# # # #         campo_pesquisa.clear()
# # # #         campo_pesquisa.send_keys(titulo)
        
# # # #         # Clica em Listar Todos
# # # #         btn_listar = driver.find_element(By.XPATH, "//a[contains(@onclick, \"CatalogingSearch.search('simple')\")]")
# # # #         btn_listar.click()
# # # #         time.sleep(2.5)
        
# # # #         # Abre o primeiro registro da lista resultante
# # # #         print(f"[LINHA {linha}] Abrindo o registro localizado...")
# # # #         btn_abrir = driver.find_element(By.XPATH, "//a[contains(@onclick, 'CatalogingSearch.openResult')]")
# # # #         btn_abrir.click()
# # # #         time.sleep(2.5)
        
# # # #         # Clica na aba Formulário
# # # #         aba_formulario = driver.find_element(By.XPATH, "//li[@data-tab='form' and contains(text(), 'Formulário')]")
# # # #         aba_formulario.click()
# # # #         time.sleep(1)
        
# # # #         # Clica no botão Editar
# # # #         btn_editar = driver.find_element(By.XPATH, "//a[contains(@onclick, 'CatalogingInput.editRecord')]")
# # # #         btn_editar.click()
# # # #         time.sleep(2)
        
# # # #         # -------------------------------------------------------------------
# # # #         # PASSO 2.3: PREENCHIMENTO DO CAMPO MARC 090 (Número de Chamada)
# # # #         # -------------------------------------------------------------------
# # # #         # 1. Classificação ($a) da div com data='090'
# # # #         if cdd:
# # # #             print(f"[LINHA {linha}] Preenchendo Classificação (090 $a) com: {cdd}")
# # # #             input_classificacao = driver.find_element(By.XPATH, "//fieldset[@data='090']//div[@data='090' and contains(@class, 'subfield')]//input[@name='a']")
# # # #             driver.execute_script("arguments[0].scrollIntoView({block:'center'});", input_classificacao)
# # # #             input_classificacao.clear()
# # # #             input_classificacao.send_keys(str(cdd))
            
# # # #         # 2. Código do Autor ($b) da div com data='090'
# # # #         if codigo_autor:
# # # #             print(f"[LINHA {linha}] Preenchendo Código do Autor (090 $b) com: {codigo_autor}")
# # # #             input_codigo_autor = driver.find_element(By.XPATH, "//fieldset[@data='090']//div[@data='090' and contains(@class, 'subfield')]//input[@name='b']")
# # # #             input_codigo_autor.clear()
# # # #             input_codigo_autor.send_keys(str(codigo_autor))
            
# # # #         driver.save_screenshot(f"linha_{linha}_02_preenchida.png")
# # # #         time.sleep(0.5)
        
# # # #         # -------------------------------------------------------------------
# # # #         # PASSO 2.4: SALVAMENTO DO REGISTRO
# # # #         # -------------------------------------------------------------------
# # # #         print(f"[LINHA {linha}] Clicando em Salvar...")
# # # #         botao_salvar = driver.find_element(By.XPATH, "//a[contains(@onclick,'CatalogingInput.saveRecord')]")
# # # #         driver.execute_script("arguments[0].scrollIntoView(true);", botao_salvar)
# # # #         time.sleep(0.5)
# # # #         driver.execute_script("arguments[0].click();", botao_salvar)
        
# # # #         time.sleep(3.5) # Tempo para processamento no banco de dados local
# # # #         driver.save_screenshot(f"linha_{linha}_03_salva.png")
# # # #         print(f"[LINHA {linha}] Registro atualizado e salvo com sucesso!")

# # # #     except Exception as erro:
# # # #         print(f"[ERRO LINHA {linha}] Ocorreu uma falha no processamento: {erro}")
# # # #         driver.save_screenshot(f"linha_{linha}_erro.png")
# # # #         # Continua para o próximo livro mesmo que a linha atual falhe
# # # #         continue

# # # # print("\n[INFO] Processo finalizado com sucesso!")
# # # # driver.quit()
# # # from selenium import webdriver
# # # from selenium.webdriver.common.by import By
# # # from selenium.webdriver.common.keys import Keys
# # # from openpyxl import load_workbook
# # # import time

# # # def tratar_nome_autor(autor_string):
# # #     """
# # #     Pega o nome do autor da planilha. Se houver mais de um (separado por vírgula, 
# # #     ponto e vírgula ou 'e'), mantém apenas o primeiro.
# # #     Inverte o nome para o formato 'Sobrenome, Nome' exigido pelo gerador Cutter.
# # #     """
# # #     if not autor_string:
# # #         return ""
    
# # #     # Isola o primeiro autor se houver múltiplos (separadores comuns)
# # #     for separador in [";", ",", " e "]:
# # #         if separador in autor_string:
# # #             autor_string = autor_string.split(separador)[0].strip()
# # #             break
            
# # #     partes = autor_string.strip().split()
# # #     if len(partes) <= 1:
# # #         return autor_string # Retorna o próprio termo se for palavra única
    
# # #     sobrenome = partes[-1]       # Último nome
# # #     nome = " ".join(partes[:-1]) # Nomes anteriores
# # #     return f"{sobrenome}, {nome}"

# # # # PLANILHA
# # # wb = load_workbook("planilha.xlsx")
# # # ws = wb["Sheet1"]

# # # # CHROME
# # # options = webdriver.ChromeOptions()
# # # driver = webdriver.Chrome(options=options)
# # # driver.maximize_window()

# # # print("[INFO] Abrindo Biblivre...")
# # # driver.get("http://localhost/Biblivre5/")
# # # time.sleep(3)

# # # # ---------------------------------------------------------------------------
# # # # PASSO 1: ROTINA DE LOGIN NO BIBLIVRE
# # # # ---------------------------------------------------------------------------
# # # try:
# # #     print("[INFO] Realizando login...")
# # #     campo_usuario = driver.find_element(By.NAME, "username")
# # #     campo_usuario.clear()
# # #     campo_usuario.send_keys("admin")
    
# # #     campo_senha = driver.find_element(By.NAME, "password")
# # #     campo_senha.clear()
# # #     campo_senha.send_keys("Bibsis2025#")
    
# # #     botao_entrar = driver.find_element(By.XPATH, "//button[contains(@onclick, \"submitForm('login'\")]")
# # #     botao_entrar.click()
    
# # #     time.sleep(3)
# # #     print("[INFO] Login realizado com sucesso!")
# # # except Exception as e:
# # #     print(f"[ERRO LOGIN] Não foi possível logar: {e}")
# # #     driver.save_screenshot("erro_login.png")
# # #     driver.quit()
# # #     exit()

# # # # Guardar o identificador da aba do Biblivre
# # # aba_biblivre = driver.current_window_handle

# # # # ---------------------------------------------------------------------------
# # # # PASSO 2: LAÇO DE REPETIÇÃO DA PLANILHA
# # # # ---------------------------------------------------------------------------
# # # for linha in range(12, 20):

# # #     print("\n" + "=" * 60)
# # #     print(f"[LINHA {linha}] Iniciando processamento")

# # #     titulo = ws[f'B{linha}'].value
# # #     subtitulo = ws[f'C{linha}'].value
# # #     autor_raw = ws[f'D{linha}'].value or ""
# # #     cdd = ws[f'K{linha}'].value or ""

# # #     if not titulo:
# # #         print(f"[LINHA {linha}] Ignorada: Título vazio.")
# # #         continue

# # #     # AJUSTE: Monta o termo de pesquisa combinando Título e Subtítulo (se houver)
# # #     titulo = str(titulo).strip()
# # #     if subtitulo:
# # #         subtitulo = str(subtitulo).strip()
# # #         termo_pesquisa = f"{titulo} {subtitulo}"
# # #     else:
# # #         termo_pesquisa = titulo

# # #     print(f"[LINHA {linha}] Livro: {titulo} | Subtítulo: {subtitulo or 'Não possui'}")
# # #     print(f"[LINHA {linha}] Termo que será pesquisado: '{termo_pesquisa}'")

# # #     # -----------------------------------------------------------------------
# # #     # PASSO 2.1: GERAR CÓDIGO CUTTER (Pegando direto do span)
# # #     # -----------------------------------------------------------------------
# # #     codigo_autor = ""
# # #     autor_formatado = tratar_nome_autor(autor_raw)

# # #     if autor_formatado:
# # #         try:
# # #             print(f"[LINHA {linha}] Consultando código Cutter para: '{autor_formatado}'")
# # #             # Abre uma nova aba para o site do Cutter
# # #             driver.execute_script("window.open('https://www.tabelacutter.com/', '_blank');")
# # #             time.sleep(2)
            
# # #             # Alterna para a aba do Cutter
# # #             abas = driver.window_handles
# # #             driver.switch_to.window(abas[-1])
            
# # #             # Digita no campo de texto do autor
# # #             input_cutter = driver.find_element(By.NAME, "e")
# # #             input_cutter.clear()
# # #             input_cutter.send_keys(autor_formatado)
            
# # #             # Clica no botão para Gerar o código
# # #             btn_gerar = driver.find_element(By.XPATH, "//button[contains(text(), 'Gerar código Cutter')]")
# # #             btn_gerar.click()
# # #             time.sleep(1.5)
            
# # #             # Captura o texto direto de dentro do <span> com o ID cutterSpan
# # #             elemento_span = driver.find_element(By.ID, "cutterSpan")
# # #             codigo_autor = elemento_span.text.strip()
            
# # #             print(f"[LINHA {linha}] Código Cutter obtido direto do span: {codigo_autor}")
            
# # #             # Fecha a aba do Cutter e retorna para a aba principal do Biblivre
# # #             driver.close()
# # #             driver.switch_to.window(aba_biblivre)
# # #             time.sleep(1)
            
# # #         except Exception as e_cutter:
# # #             print(f"[AVISO CUTTER] Falha ao gerar código Cutter na linha {linha}: {e_cutter}")
# # #             try:
# # #                 driver.switch_to.window(aba_biblivre)
# # #             except:
# # #                 pass

# # #     # -----------------------------------------------------------------------
# # #     # PASSO 2.2: PESQUISAR E EDITAR O REGISTRO NO BIBLIVRE
# # #     # -----------------------------------------------------------------------
# # #     try:
# # #         print(f"[LINHA {linha}] Navegando até Catalogação -> Bibliográfica...")
# # #         menu_catalogacao = driver.find_element(By.CSS_SELECTOR, "li.menu_cataloging")
# # #         menu_catalogacao.click()
# # #         time.sleep(0.5)
        
# # #         submenu_bibliografica = driver.find_element(By.XPATH, "//li[@data-action='cataloging_bibliographic']/a")
# # #         submenu_bibliografica.click()
# # #         time.sleep(2.5)

# # #         # AJUSTE: Agora envia o termo combinado (Título + Subtítulo) para o campo de pesquisa
# # #         print(f"[LINHA {linha}] Pesquisando no Biblivre por: '{termo_pesquisa}'")
# # #         campo_pesquisa = driver.find_element(By.NAME, "query")
# # #         campo_pesquisa.clear()
# # #         campo_pesquisa.send_keys(termo_pesquisa)
        
# # #         # Clica em Listar Todos
# # #         btn_listar = driver.find_element(By.XPATH, "//a[contains(@onclick, \"CatalogingSearch.search('simple')\")]")
# # #         btn_listar.click()
# # #         time.sleep(2.5)
        
# # #         # Abre o primeiro registro da lista resultante
# # #         print(f"[LINHA {linha}] Abrindo o registro localizado...")
# # #         btn_abrir = driver.find_element(By.XPATH, "//a[contains(@onclick, 'CatalogingSearch.openResult')]")
# # #         btn_abrir.click()
# # #         time.sleep(2.5)
        
# # #         # Clica na aba Formulário
# # #         aba_formulario = driver.find_element(By.XPATH, "//li[@data-tab='form' and contains(text(), 'Formulário')]")
# # #         aba_formulario.click()
# # #         time.sleep(1)
        
# # #         # Clica no botão Editar
# # #         btn_editar = driver.find_element(By.XPATH, "//a[contains(@onclick, 'CatalogingInput.editRecord')]")
# # #         btn_editar.click()
# # #         time.sleep(2)
        
# # #         # -------------------------------------------------------------------
# # #         # PASSO 2.3: PREENCHIMENTO DO CAMPO MARC 090 (Número de Chamada)
# # #         # -------------------------------------------------------------------
# # #         if cdd:
# # #             print(f"[LINHA {linha}] Preenchendo Classificação (090 $a) com: {cdd}")
# # #             input_classificacao = driver.find_element(By.XPATH, "//fieldset[@data='090']//div[@data='090' and contains(@class, 'subfield')]//input[@name='a']")
# # #             driver.execute_script("arguments[0].scrollIntoView({block:'center'});", input_classificacao)
# # #             input_classificacao.clear()
# # #             input_classificacao.send_keys(str(cdd))
            
# # #         if codigo_autor:
# # #             print(f"[LINHA {linha}] Preenchendo Código do Autor (090 $b) com: {codigo_autor}")
# # #             input_codigo_autor = driver.find_element(By.XPATH, "//fieldset[@data='090']//div[@data='090' and contains(@class, 'subfield')]//input[@name='b']")
# # #             input_codigo_autor.clear()
# # #             input_codigo_autor.send_keys(str(codigo_autor))
            
# # #         driver.save_screenshot(f"linha_{linha}_02_preenchida.png")
# # #         time.sleep(0.5)
        
# # #         # -------------------------------------------------------------------
# # #         # PASSO 2.4: SALVAMENTO DO REGISTRO
# # #         # -------------------------------------------------------------------
# # #         print(f"[LINHA {linha}] Clicando em Salvar...")
# # #         botao_salvar = driver.find_element(By.XPATH, "//a[contains(@onclick,'CatalogingInput.saveRecord')]")
# # #         driver.execute_script("arguments[0].scrollIntoView(true);", botao_salvar)
# # #         time.sleep(0.5)
# # #         driver.execute_script("arguments[0].click();", botao_salvar)
        
# # #         time.sleep(3.5)
# # #         driver.save_screenshot(f"linha_{linha}_03_salva.png")
# # #         print(f"[LINHA {linha}] Registro atualizado e salvo com sucesso!")

# # #     except Exception as erro:
# # #         print(f"[ERRO LINHA {linha}] Ocorreu uma falha no processamento: {erro}")
# # #         driver.save_screenshot(f"linha_{linha}_erro.png")
# # #         continue

# # # print("\n[INFO] Processo finalizado com sucesso!")
# # # driver.quit()
# # from selenium import webdriver
# # from selenium.webdriver.common.by import By
# # from selenium.webdriver.common.keys import Keys
# # from openpyxl import load_workbook
# # import time

# # def tratar_nome_autor(autor_string):
# #     """
# #     Pega o nome do autor da planilha. Se houver mais de um (separado por vírgula, 
# #     ponto e vírgula ou 'e'), mantém apenas o primeiro.
# #     Inverte o nome para o formato 'Sobrenome, Nome' exigido pelo gerador Cutter.
# #     """
# #     if not autor_string:
# #         return ""
    
# #     # Isola o primeiro autor se houver múltiplos (separadores comuns)
# #     for separador in [";", ",", " e "]:
# #         if separador in autor_string:
# #             autor_string = autor_string.split(separador)[0].strip()
# #             break
            
# #     partes = autor_string.strip().split()
# #     if len(partes) <= 1:
# #         return autor_string # Retorna o próprio termo se for palavra única
    
# #     sobrenome = partes[-1]       # Último nome
# #     nome = " ".join(partes[:-1]) # Nomes anteriores
# #     return f"{sobrenome}, {nome}"

# # # PLANILHA
# # wb = load_workbook("planilha.xlsx")
# # ws = wb["Sheet1"]

# # # LISTA PARA GUARDAR AS OBRAS NÃO INDEXADAS (DIVERGÊNCIA DE AUTOR)
# # obras_nao_indexadas = []

# # # CHROME
# # options = webdriver.ChromeOptions()
# # driver = webdriver.Chrome(options=options)
# # driver.maximize_window()

# # print("[INFO] Abrindo Biblivre...")
# # driver.get("http://localhost/Biblivre5/")
# # time.sleep(3)

# # # ---------------------------------------------------------------------------
# # # PASSO 1: ROTINA DE LOGIN NO BIBLIVRE
# # # ---------------------------------------------------------------------------
# # try:
# #     print("[INFO] Realizando login...")
# #     campo_usuario = driver.find_element(By.NAME, "username")
# #     campo_usuario.clear()
# #     campo_usuario.send_keys("admin")
    
# #     campo_senha = driver.find_element(By.NAME, "password")
# #     campo_senha.clear()
# #     campo_senha.send_keys("Bibsis2025#")
    
# #     botao_entrar = driver.find_element(By.XPATH, "//button[contains(@onclick, \"submitForm('login'\")]")
# #     botao_entrar.click()
    
# #     time.sleep(3)
# #     print("[INFO] Login realizado com sucesso!")
# # except Exception as e:
# #     print(f"[ERRO LOGIN] Não foi possível logar: {e}")
# #     driver.save_screenshot("erro_login.png")
# #     driver.quit()
# #     exit()

# # # Guardar o identificador da aba do Biblivre
# # aba_biblivre = driver.current_window_handle

# # # ---------------------------------------------------------------------------
# # # PASSO 2: LAÇO DE REPETIÇÃO DA PLANILHA
# # # ---------------------------------------------------------------------------
# # for linha in range(17, 20):

# #     print("\n" + "=" * 60)
# #     print(f"[LINHA {linha}] Iniciando processamento")

# #     titulo = ws[f'B{linha}'].value
# #     subtitulo = ws[f'C{linha}'].value
# #     autor_raw = ws[f'D{linha}'].value or ""
# #     cdd = ws[f'K{linha}'].value or ""

# #     if not titulo:
# #         print(f"[LINHA {linha}] Ignorada: Título vazio.")
# #         continue

# #     titulo = str(titulo).strip()
# #     subtitulo = str(subtitulo).strip() if subtitulo else ""
# #     termo_pesquisa = f"{titulo} {subtitulo}".strip()

# #     print(f"[LINHA {linha}] Livro: {titulo} | Subtítulo: {subtitulo or 'Não possui'}")
# #     print(f"[LINHA {linha}] Autor Esperado (Planilha): '{autor_raw}'")

# #     # -----------------------------------------------------------------------
# #     # PASSO 2.1: GERAR CÓDIGO CUTTER (Pegando direto do span)
# #     # -----------------------------------------------------------------------
# #     codigo_autor = ""
# #     autor_formatado = tratar_nome_autor(autor_raw)

# #     if autor_formatado:
# #         try:
# #             print(f"[LINHA {linha}] Consultando código Cutter para: '{autor_formatado}'")
# #             driver.execute_script("window.open('https://www.tabelacutter.com/', '_blank');")
# #             time.sleep(2)
            
# #             abas = driver.window_handles
# #             driver.switch_to.window(abas[-1])
            
# #             input_cutter = driver.find_element(By.NAME, "e")
# #             input_cutter.clear()
# #             input_cutter.send_keys(autor_formatado)
            
# #             btn_gerar = driver.find_element(By.XPATH, "//button[contains(text(), 'Gerar código Cutter')]")
# #             btn_gerar.click()
# #             time.sleep(1.5)
            
# #             elemento_span = driver.find_element(By.ID, "cutterSpan")
# #             codigo_autor = elemento_span.text.strip()
            
# #             print(f"[LINHA {linha}] Código Cutter obtido direto do span: {codigo_autor}")
            
# #             driver.close()
# #             driver.switch_to.window(aba_biblivre)
# #             time.sleep(1)
            
# #         except Exception as e_cutter:
# #             print(f"[AVISO CUTTER] Falha ao gerar código Cutter na linha {linha}: {e_cutter}")
# #             try:
# #                 driver.switch_to.window(aba_biblivre)
# #             except:
# #                 pass

# #     # -----------------------------------------------------------------------
# #     # PASSO 2.2: PESQUISAR E EDITAR O REGISTRO NO BIBLIVRE COM VALIDAÇÃO
# #     # -----------------------------------------------------------------------
# #     try:
# #         print(f"[LINHA {linha}] Navegando até Catalogação -> Bibliográfica...")
# #         menu_catalogacao = driver.find_element(By.CSS_SELECTOR, "li.menu_cataloging")
# #         menu_catalogacao.click()
# #         time.sleep(0.5)
        
# #         submenu_bibliografica = driver.find_element(By.XPATH, "//li[@data-action='cataloging_bibliographic']/a")
# #         submenu_bibliografica.click()
# #         time.sleep(2.5)

# #         print(f"[LINHA {linha}] Pesquisando no Biblivre por: '{termo_pesquisa}'")
# #         campo_pesquisa = driver.find_element(By.NAME, "query")
# #         campo_pesquisa.clear()
# #         campo_pesquisa.send_keys(termo_pesquisa)
        
# #         btn_listar = driver.find_element(By.XPATH, "//a[contains(@onclick, \"CatalogingSearch.search('simple')\")]")
# #         btn_listar.click()
# #         time.sleep(2.5)
        
# #         print(f"[LINHA {linha}] Abrindo o registro localizado...")
# #         btn_abrir = driver.find_element(By.XPATH, "//a[contains(@onclick, 'CatalogingSearch.openResult')]")
# #         btn_abrir.click()
# #         time.sleep(2.5)
        
# #         aba_formulario = driver.find_element(By.XPATH, "//li[@data-tab='form' and contains(text(), 'Formulário')]")
# #         aba_formulario.click()
# #         time.sleep(1)
        
# #         btn_editar = driver.find_element(By.XPATH, "//a[contains(@onclick, 'CatalogingInput.editRecord')]")
# #         btn_editar.click()
# #         time.sleep(2)
        
# #         # -------------------------------------------------------------------
# #         # AJUSTE: VERIFICAÇÃO SE O AUTOR DO SISTEMA BATE COM O DA PLANILHA
# #         # -------------------------------------------------------------------
# #         # Localiza o input específico de autoridades (autor principal)
# #         input_autor_sistema = driver.find_element(By.XPATH, "//input[@data-ac='authorities' and @name='a']")
# #         autor_sistema_texto = input_autor_sistema.get_attribute("value").strip()
        
# #         print(f"[LINHA {linha}] Autor encontrado no Biblivre: '{autor_sistema_texto}'")
        
# #         # Compara de forma tolerante (removendo espaços extras e ignorando maiúsculas/minúsculas)
# #         if autor_raw.strip().lower() != autor_sistema_texto.lower():
# #             print(f"[ALERTA LINHA {linha}] Divergência detectada! Planilha: '{autor_raw}' vs Sistema: '{autor_sistema_texto}'")
            
# #             # Salva os dados da obra na lista de não indexadas
# #             obras_nao_indexadas.append({
# #                 "linha": linha,
# #                 "titulo": titulo,
# #                 "subtitulo": subtitulo,
# #                 "autor_planilha": autor_raw,
# #                 "autor_sistema": autor_sistema_texto
# #             })
            
# #             # Clica em Cancelar para sair do modo de edição com segurança
# #             try:
# #                 btn_cancelar = driver.find_element(By.XPATH, "//a[contains(@onclick, 'CatalogingInput.cancelEdit')]")
# #                 driver.execute_script("arguments[0].click();", btn_cancelar)
# #                 time.sleep(1.5)
# #             except:
# #                 pass
                
# #             print(f"[LINHA {linha}] Pulando para a próxima obra sem alterar nada.")
# #             continue # Interrompe o fluxo desta linha aqui e vai para o próximo livro
            
# #         # -------------------------------------------------------------------
# #         # PASSO 2.3: PREENCHIMENTO DO CAMPO MARC 090 (Número de Chamada)
# #         # -------------------------------------------------------------------
# #         if cdd:
# #             print(f"[LINHA {linha}] Preenchendo Classificação (090 $a) com: {cdd}")
# #             input_classificacao = driver.find_element(By.XPATH, "//fieldset[@data='090']//div[@data='090' and contains(@class, 'subfield')]//input[@name='a']")
# #             driver.execute_script("arguments[0].scrollIntoView({block:'center'});", input_classificacao)
# #             input_classificacao.clear()
# #             input_classificacao.send_keys(str(cdd))
            
# #         if codigo_autor:
# #             print(f"[LINHA {linha}] Preenchendo Código do Autor (090 $b) com: {codigo_autor}")
# #             input_codigo_autor = driver.find_element(By.XPATH, "//fieldset[@data='090']//div[@data='090' and contains(@class, 'subfield')]//input[@name='b']")
# #             input_codigo_autor.clear()
# #             input_codigo_autor.send_keys(str(codigo_autor))
            
# #         driver.save_screenshot(f"linha_{linha}_02_preenchida.png")
# #         time.sleep(0.5)
        
# #         # -------------------------------------------------------------------
# #         # PASSO 2.4: SALVAMENTO DO REGISTRO
# #         # -------------------------------------------------------------------
# #         print(f"[LINHA {linha}] Clicando em Salvar...")
# #         botao_salvar = driver.find_element(By.XPATH, "//a[contains(@onclick,'CatalogingInput.saveRecord')]")
# #         driver.execute_script("arguments[0].scrollIntoView(true);", botao_salvar)
# #         time.sleep(0.5)
# #         driver.execute_script("arguments[0].click();", botao_salvar)
        
# #         time.sleep(3.5)
# #         driver.save_screenshot(f"linha_{linha}_03_salva.png")
# #         print(f"[LINHA {linha}] Registro atualizado e salvo com sucesso!")

# #     except Exception as erro:
# #         print(f"[ERRO LINHA {linha}] Ocorreu uma falha no processamento: {erro}")
# #         driver.save_screenshot(f"linha_{linha}_erro.png")
# #         continue

# # print("\n" + "=" * 60)
# # print("[INFO] Processo de processamento finalizado!")
# # driver.quit()

# # # ---------------------------------------------------------------------------
# # # PASSO 3: EXIBIÇÃO DO RELATÓRIO DE OBRAS NÃO INDEXADAS
# # # ---------------------------------------------------------------------------
# # print("\n" + "X" * 60)
# # print(f"   RELATÓRIO DE OBRAS NÃO INDEXADAS (TOTAL: {len(obras_nao_indexadas)})")
# # print("X" * 60)

# # if obras_nao_indexadas:
# #     for obra in obras_nao_indexadas:
# #         print(f"\n-> Linha Planilha: {obra['linha']}")
# #         print(f"   Título: {obra['titulo']}")
# #         if obra['subtitulo']:
# #             print(f"   Subtítulo: {obra['subtitulo']}")
# #         print(f"   Autor na Planilha: {obra['autor_planilha']}")
# #         print(f"   Autor no Biblivre: {obra['autor_sistema']}")
# # else:
# #     print("\n[SUCESSO] Todas as obras pesquisadas bateram com os autores e foram indexadas!")
# # print("\n" + "X" * 60)
# from selenium import webdriver
# from selenium.webdriver.common.by import By
# from selenium.webdriver.common.keys import Keys
# from openpyxl import load_workbook
# import time

# def tratar_nome_autor(autor_string):
#     """
#     Pega o nome do autor da planilha. Se houver mais de um (separado por vírgula, 
#     ponto e vírgula ou 'e'), mantém apenas o primeiro.
#     Inverte o nome para o formato 'Sobrenome, Nome' exigido pelo gerador Cutter.
#     """
#     if not autor_string:
#         return ""
    
#     # Isola o primeiro autor se houver múltiplos (separadores comuns)
#     for separador in [";", ",", " e "]:
#         if separador in autor_string:
#             autor_string = autor_string.split(separador)[0].strip()
#             break
            
#     partes = autor_string.strip().split()
#     if len(partes) <= 1:
#         return autor_string # Retorna o próprio termo se for palavra única
    
#     sobrenome = partes[-1]       # Último nome
#     nome = " ".join(partes[:-1]) # Nomes anteriores
#     return f"{sobrenome}, {nome}"

# # PLANILHA
# wb = load_workbook("planilha.xlsx")
# ws = wb["Sheet1"]

# # LISTA PARA GUARDAR AS OBRAS NÃO INDEXADAS (DIVERGÊNCIA DE AUTOR)
# obras_nao_indexadas = []

# # CHROME
# options = webdriver.ChromeOptions()
# driver = webdriver.Chrome(options=options)
# driver.maximize_window()

# print("[INFO] Abrindo Biblivre...")
# driver.get("http://localhost/Biblivre5/")
# time.sleep(3)

# # ---------------------------------------------------------------------------
# # PASSO 1: ROTINA DE LOGIN NO BIBLIVRE
# # ---------------------------------------------------------------------------
# try:
#     print("[INFO] Realizando login...")
#     campo_usuario = driver.find_element(By.NAME, "username")
#     campo_usuario.clear()
#     campo_usuario.send_keys("admin")
    
#     campo_senha = driver.find_element(By.NAME, "password")
#     campo_senha.clear()
#     campo_senha.send_keys("Bibsis2025#")
    
#     botao_entrar = driver.find_element(By.XPATH, "//button[contains(@onclick, \"submitForm('login'\")]")
#     botao_entrar.click()
    
#     time.sleep(3)
#     print("[INFO] Login realizado com sucesso!")
# except Exception as e:
#     print(f"[ERRO LOGIN] Não foi possível logar: {e}")
#     driver.save_screenshot("erro_login.png")
#     driver.quit()
#     exit()

# # Guardar o identificador da aba do Biblivre
# aba_biblivre = driver.current_window_handle

# # ---------------------------------------------------------------------------
# # PASSO 2: LAÇO DE REPETIÇÃO DA PLANILHA
# # ---------------------------------------------------------------------------
# for linha in range(22, 72):

#     print("\n" + "=" * 60)
#     print(f"[LINHA {linha}] Iniciando processamento")

#     titulo = ws[f'B{linha}'].value
#     subtitulo = ws[f'C{linha}'].value
#     autor_raw = ws[f'D{linha}'].value or ""
#     cdd = ws[f'K{linha}'].value or ""

#     if not titulo:
#         print(f"[LINHA {linha}] Ignorada: Título vazio.")
#         continue

#     titulo = str(titulo).strip()
#     subtitulo = str(subtitulo).strip() if subtitulo else ""
#     termo_pesquisa = f"{titulo} {subtitulo}".strip()

#     print(f"[LINHA {linha}] Livro: {titulo} | Subtítulo: {subtitulo or 'Não possui'}")
#     print(f"[LINHA {linha}] Autor Esperado (Planilha): '{autor_raw}'")

#     # -----------------------------------------------------------------------
#     # PASSO 2.1: GERAR CÓDIGO CUTTER (Pegando direto do span)
#     # -----------------------------------------------------------------------
#     codigo_autor = ""
#     autor_formatado = tratar_nome_autor(autor_raw)

#     if autor_formatado:
#         try:
#             print(f"[LINHA {linha}] Consultando código Cutter para: '{autor_formatado}'")
#             driver.execute_script("window.open('https://www.tabelacutter.com/', '_blank');")
#             time.sleep(2)
            
#             abas = driver.window_handles
#             driver.switch_to.window(abas[-1])
            
#             input_cutter = driver.find_element(By.NAME, "e")
#             input_cutter.clear()
#             input_cutter.send_keys(autor_formatado)
            
#             btn_gerar = driver.find_element(By.XPATH, "//button[contains(text(), 'Gerar código Cutter')]")
#             btn_gerar.click()
#             time.sleep(1.5)
            
#             elemento_span = driver.find_element(By.ID, "cutterSpan")
#             codigo_autor = elemento_span.text.strip()
            
#             print(f"[LINHA {linha}] Código Cutter obtido direto do span: {codigo_autor}")
            
#             driver.close()
#             driver.switch_to.window(aba_biblivre)
#             time.sleep(1)
            
#         except Exception as e_cutter:
#             print(f"[AVISO CUTTER] Falha ao gerar código Cutter na linha {linha}: {e_cutter}")
#             try:
#                 driver.switch_to.window(aba_biblivre)
#             except:
#                 pass

#     # -----------------------------------------------------------------------
#     # PASSO 2.2: PESQUISAR E EDITAR O REGISTRO NO BIBLIVRE COM VALIDAÇÃO
#     # -----------------------------------------------------------------------
#     try:
#         print(f"[LINHA {linha}] Navegando até Catalogação -> Bibliográfica...")
#         menu_catalogacao = driver.find_element(By.CSS_SELECTOR, "li.menu_cataloging")
#         menu_catalogacao.click()
#         time.sleep(0.5)
        
#         submenu_bibliografica = driver.find_element(By.XPATH, "//li[@data-action='cataloging_bibliographic']/a")
#         submenu_bibliografica.click()
#         time.sleep(2.5)

#         print(f"[LINHA {linha}] Pesquisando no Biblivre por: '{termo_pesquisa}'")
#         campo_pesquisa = driver.find_element(By.NAME, "query")
#         campo_pesquisa.clear()
#         campo_pesquisa.send_keys(termo_pesquisa)
        
#         btn_listar = driver.find_element(By.XPATH, "//a[contains(@onclick, \"CatalogingSearch.search('simple')\")]")
#         btn_listar.click()
#         time.sleep(2.5)
        
#         print(f"[LINHA {linha}] Abrindo o registro localizado...")
#         btn_abrir = driver.find_element(By.XPATH, "//a[contains(@onclick, 'CatalogingSearch.openResult')]")
#         btn_abrir.click()
#         time.sleep(2.5)
        
#         aba_formulario = driver.find_element(By.XPATH, "//li[@data-tab='form' and contains(text(), 'Formulário')]")
#         aba_formulario.click()
#         time.sleep(1)
        
#         btn_editar = driver.find_element(By.XPATH, "//a[contains(@onclick, 'CatalogingInput.editRecord')]")
#         btn_editar.click()
#         time.sleep(2)
        
#         # Localiza o input específico de autoridades (autor principal)
#         input_autor_sistema = driver.find_element(By.XPATH, "//input[@data-ac='authorities' and @name='a']")
#         autor_sistema_texto = input_autor_sistema.get_attribute("value").strip()
        
#         print(f"[LINHA {linha}] Autor encontrado no Biblivre: '{autor_sistema_texto}'")
        
#         # -------------------------------------------------------------------
#         # VERIFICAÇÃO SE O AUTOR DO SISTEMA BATE COM O DA PLANILHA
#         # -------------------------------------------------------------------
#         if autor_raw.strip().lower() != autor_sistema_texto.lower():
#             print(f"[ALERTA LINHA {linha}] Divergência detectada! Planilha: '{autor_raw}' vs Sistema: '{autor_sistema_texto}'")
            
#             # Salva os dados da obra na lista de não indexadas
#             obras_nao_indexadas.append({
#                 "linha": linha,
#                 "titulo": titulo,
#                 "subtitulo": subtitulo,
#                 "autor_planilha": autor_raw,
#                 "autor_sistema": autor_sistema_texto
#             })
            
#             # NOVO AJUSTE: Rotina completa para cancelar a edição confirmando no "Sim"
#             try:
#                 print(f"[LINHA {linha}] Clicando em Cancelar...")
#                 btn_cancelar = driver.find_element(By.XPATH, "//a[contains(@onclick, 'CatalogingInput.cancelEdit')]")
#                 driver.execute_script("arguments[0].click();", btn_cancelar)
#                 time.sleep(1.0) # Espera a caixinha de confirmação aparecer
                
#                 print(f"[LINHA {linha}] Confirmando cancelamento clicando em 'Sim'...")
#                 btn_sim = driver.find_element(By.XPATH, "//a[contains(@class, 'button') and text()='Sim']")
#                 driver.execute_script("arguments[0].click();", btn_sim)
#                 time.sleep(1.5) # Aguarda fechar o formulário completamente
#             except Exception as e_cancelar:
#                 print(f"[AVISO LINHA {linha}] Falha ao tentar fechar caixinha de confirmação: {e_cancelar}")
                
#             print(f"[LINHA {linha}] Avançando de linha com segurança.")
#             continue
            
#         # -------------------------------------------------------------------
#         # PASSO 2.3: PREENCHIMENTO DO CAMPO MARC 090 (Número de Chamada)
#         # -------------------------------------------------------------------
#         if cdd:
#             print(f"[LINHA {linha}] Preenchendo Classificação (090 $a) com: {cdd}")
#             input_classificacao = driver.find_element(By.XPATH, "//fieldset[@data='090']//div[@data='090' and contains(@class, 'subfield')]//input[@name='a']")
#             driver.execute_script("arguments[0].scrollIntoView({block:'center'});", input_classificacao)
#             input_classificacao.clear()
#             input_classificacao.send_keys(str(cdd))
            
#         if codigo_autor:
#             print(f"[LINHA {linha}] Preenchendo Código do Autor (090 $b) com: {codigo_autor}")
#             input_codigo_autor = driver.find_element(By.XPATH, "//fieldset[@data='090']//div[@data='090' and contains(@class, 'subfield')]//input[@name='b']")
#             input_codigo_autor.clear()
#             input_codigo_autor.send_keys(str(codigo_autor))
            
#         driver.save_screenshot(f"linha_{linha}_02_preenchida.png")
#         time.sleep(0.5)
        
#         # -------------------------------------------------------------------
#         # PASSO 2.4: SALVAMENTO DO REGISTRO
#         # -------------------------------------------------------------------
#         print(f"[LINHA {linha}] Clicando em Salvar...")
#         botao_salvar = driver.find_element(By.XPATH, "//a[contains(@onclick,'CatalogingInput.saveRecord')]")
#         driver.execute_script("arguments[0].scrollIntoView(true);", botao_salvar)
#         time.sleep(0.5)
#         driver.execute_script("arguments[0].click();", botao_salvar)
        
#         time.sleep(3.5)
#         driver.save_screenshot(f"linha_{linha}_03_salva.png")
#         print(f"[LINHA {linha}] Registro atualizado e salvo com sucesso!")

#     except Exception as erro:
#         print(f"[ERRO LINHA {linha}] Ocorreu uma falha no processamento: {erro}")
#         driver.save_screenshot(f"linha_{linha}_erro.png")
#         continue

# print("\n" + "=" * 60)
# print("[INFO] Processo de processamento finalizado!")
# driver.quit()

# # ---------------------------------------------------------------------------
# # PASSO 3: EXIBIÇÃO DO RELATÓRIO DE OBRAS NÃO INDEXADAS
# # ---------------------------------------------------------------------------
# print("\n" + "X" * 60)
# print(f"   RELATÓRIO DE OBRAS NÃO INDEXADAS (TOTAL: {len(obras_nao_indexadas)})")
# print("X" * 60)

# if obras_nao_indexadas:
#     for obra in obras_nao_indexadas:
#         print(f"\n-> Linha Planilha: {obra['linha']}")
#         print(f"   Título: {obra['titulo']}")
#         if obra['subtitulo']:
#             print(f"   Subtítulo: {obra['subtitulo']}")
#         print(f"   Autor na Planilha: {obra['autor_planilha']}")
#         print(f"   Autor no Biblivre: {obra['autor_sistema']}")
# else:
#     print("\n[SUCESSO] Todas as obras pesquisadas bateram com os autores e foram indexadas!")
# print("\n" + "X" * 60)

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from openpyxl import load_workbook
import time

def tratar_nome_autor(autor_string):
    """
    Pega o nome do autor da planilha. Se houver mais de um (separado por vírgula, 
    ponto e vírgula ou 'e'), mantém apenas o primeiro.
    Inverte o nome para o formato 'Sobrenome, Nome' exigido pelo gerador Cutter.
    """
    if not autor_string:
        return ""
    
    # Isola o primeiro autor se houver múltiplos (separadores comuns)
    for separador in [";", ",", " e "]:
        if separador in autor_string:
            autor_string = autor_string.split(separador)[0].strip()
            break
            
    partes = autor_string.strip().split()
    if len(partes) <= 1:
        return autor_string # Retorna o próprio termo se for palavra única
    
    sobrenome = partes[-1]       # Último nome
    nome = " ".join(partes[:-1]) # Nomes anteriores
    return f"{sobrenome}, {nome}"

def obter_inicial_titulo(titulo_string):
    """
    Remove artigos do início do título e retorna a primeira letra
    da primeira palavra significativa em formato minúsculo.
    """
    if not titulo_string:
        return ""
    
    # Lista de artigos comuns no início de títulos (maiúsculas e minúsculas)
    artigos = ["o", "a", "os", "as", "um", "uma", "uns", "umas", "de", "do", "da"]
    
    # Divide o título em palavras limpando espaços extras
    palavras = titulo_string.strip().split()
    
    if not palavras:
        return ""
        
    # Verifica se a primeira palavra é um artigo/preposição irrelevante
    while palavras and palavras[0].lower() in artigos:
        palavras.pop(0) # Remove o artigo do início
        
    if palavras:
        # Pega a primeira palavra restante, remove acentos/caracteres se necessário e extrai a primeira letra em minúsculo
        primeira_palavra = palavras[0]
        # Garante o retorno apenas da primeira letra minúscula (ex: 'h' de 'historias')
        for caractere in primeira_palavra:
            if caractere.isalpha():
                return caractere.lower()
                
    # Caso o título acabe ficando vazio após a limpeza, pega a inicial do título original por segurança
    return titulo_string.strip()[0].lower()


# PLANILHA
# wb = load_workbook("planilha.xlsx")
# wb = load_workbook("planilha (1).xlsx")
wb = load_workbook("Planilha_Catalogacao (1).xlsx")


ws = wb["Sheet1"]

# LISTA PARA GUARDAR AS OBRAS NÃO INDEXADAS (DIVERGÊNCIA DE AUTOR)
obras_nao_indexadas = []

# CHROME
options = webdriver.ChromeOptions()
driver = webdriver.Chrome(options=options)
driver.maximize_window()

print("[INFO] Abrindo Biblivre...")
driver.get("http://localhost/Biblivre5/")
time.sleep(3)

# ---------------------------------------------------------------------------
# PASSO 1: ROTINA DE LOGIN NO BIBLIVRE
# ---------------------------------------------------------------------------
try:
    print("[INFO] Realizando login...")
    campo_usuario = driver.find_element(By.NAME, "username")
    campo_usuario.clear()
    campo_usuario.send_keys("admin")
    
    campo_senha = driver.find_element(By.NAME, "password")
    campo_senha.clear()
    campo_senha.send_keys("Bibsis2025#")
    
    botao_entrar = driver.find_element(By.XPATH, "//button[contains(@onclick, \"submitForm('login'\")]")
    botao_entrar.click()
    
    time.sleep(3)
    print("[INFO] Login realizado com sucesso!")
except Exception as e:
    print(f"[ERRO LOGIN] Não foi possível logar: {e}")
    driver.save_screenshot("erro_login.png")
    driver.quit()
    exit()

# Guardar o identificador da aba do Biblivre
aba_biblivre = driver.current_window_handle

# ---------------------------------------------------------------------------
# PASSO 2: LAÇO DE REPETIÇÃO DA PLANILHA
# ---------------------------------------------------------------------------
# Linhas corretas
# for linha in range(59, 72):
# for linha in range(73, 80):
# Fiz wessa debaixo
# for linha in range(81, 90):
# Precisa refazer do inicio ate ai e dai em diante.
# for linha in range(2, 20):
# for linha in range(22, 72):
# for linha in range(73, 80):
# for linha in range(81, 90):
# for linha in range(91, 96):
# for linha in range(99, 102):
# for linha in range(106, 107):
# for linha in range(108, 110):
# for linha in range(117, 300):
# for linha in range(300, 350):
# for linha in range(350, 400):
# for linha in range(400, 450):
# for linha in range(450, 1322):
# for linha in range(1322,1355):
for linha in range(1355,1388):


    print("\n" + "=" * 60)
    print(f"[LINHA {linha}] Iniciando processamento")

    titulo = ws[f'B{linha}'].value
    subtitulo = ws[f'C{linha}'].value
    autor_raw = ws[f'D{linha}'].value or ""
    cdd = ws[f'K{linha}'].value or ""

    if not titulo:
        print(f"[LINHA {linha}] Ignorada: Título vazio.")
        continue

    titulo = str(titulo).strip()
    subtitulo = str(subtitulo).strip() if subtitulo else ""
    termo_pesquisa = f"{titulo} {subtitulo}".strip()

    print(f"[LINHA {linha}] Livro: {titulo} | Subtítulo: {subtitulo or 'Não possui'}")
    print(f"[LINHA {linha}] Autor Esperado (Planilha): '{autor_raw}'")

    # -----------------------------------------------------------------------
    # PASSO 2.1: GERAR CÓDIGO CUTTER (Pegando direto do span)
    # -----------------------------------------------------------------------
    codigo_autor = ""
    autor_formatado = tratar_nome_autor(autor_raw)

    if autor_formatado:
        try:
            print(f"[LINHA {linha}] Consultando código Cutter para: '{autor_formatado}'")
            driver.execute_script("window.open('https://www.tabelacutter.com/', '_blank');")
            time.sleep(2)
            
            abas = driver.window_handles
            driver.switch_to.window(abas[-1])
            
            input_cutter = driver.find_element(By.NAME, "e")
            input_cutter.clear()
            input_cutter.send_keys(autor_formatado)
            
            btn_gerar = driver.find_element(By.XPATH, "//button[contains(text(), 'Gerar código Cutter')]")
            btn_gerar.click()
            time.sleep(1.5)
            
            elemento_span = driver.find_element(By.ID, "cutterSpan")
            codigo_base = elemento_span.text.strip()
            
            # NOVO AJUSTE: Junta o código gerado com a inicial minúscula do título limpo
            inicial_titulo = obter_inicial_titulo(titulo)
            codigo_autor = f"{codigo_base}{inicial_titulo}"
            
            print(f"[LINHA {linha}] Código Cutter Base: {codigo_base} | Inicial do título: '{inicial_titulo}' -> Código Final: {codigo_autor}")
            
            driver.close()
            driver.switch_to.window(aba_biblivre)
            time.sleep(1)
            
        except Exception as e_cutter:
            print(f"[AVISO CUTTER] Falha ao gerar código Cutter na linha {linha}: {e_cutter}")
            try:
                driver.switch_to.window(aba_biblivre)
            except:
                pass

    # -----------------------------------------------------------------------
    # PASSO 2.2: PESQUISAR E EDITAR O REGISTRO NO BIBLIVRE COM VALIDAÇÃO
    # -----------------------------------------------------------------------
    try:
        print(f"[LINHA {linha}] Navegando até Catalogação -> Bibliográfica...")
        menu_catalogacao = driver.find_element(By.CSS_SELECTOR, "li.menu_cataloging")
        menu_catalogacao.click()
        time.sleep(0.5)
        
        submenu_bibliografica = driver.find_element(By.XPATH, "//li[@data-action='cataloging_bibliographic']/a")
        submenu_bibliografica.click()
        time.sleep(2.5)

        print(f"[LINHA {linha}] Pesquisando no Biblivre por: '{termo_pesquisa}'")
        campo_pesquisa = driver.find_element(By.NAME, "query")
        campo_pesquisa.clear()
        campo_pesquisa.send_keys(termo_pesquisa)
        
        btn_listar = driver.find_element(By.XPATH, "//a[contains(@onclick, \"CatalogingSearch.search('simple')\")]")
        btn_listar.click()
        time.sleep(2.5)
        
        print(f"[LINHA {linha}] Abrindo o registro localizado...")
        btn_abrir = driver.find_element(By.XPATH, "//a[contains(@onclick, 'CatalogingSearch.openResult')]")
        btn_abrir.click()
        time.sleep(2.5)
        
        aba_formulario = driver.find_element(By.XPATH, "//li[@data-tab='form' and contains(text(), 'Formulário')]")
        aba_formulario.click()
        time.sleep(1)
        
        btn_editar = driver.find_element(By.XPATH, "//a[contains(@onclick, 'CatalogingInput.editRecord')]")
        btn_editar.click()
        time.sleep(2)
        
        # Localiza o input específico de autoridades (autor principal)
        input_autor_sistema = driver.find_element(By.XPATH, "//input[@data-ac='authorities' and @name='a']")
        autor_sistema_texto = input_autor_sistema.get_attribute("value").strip()
        
        print(f"[LINHA {linha}] Autor encontrado no Biblivre: '{autor_sistema_texto}'")
        
        # VERIFICAÇÃO SE O AUTOR DO SISTEMA BATE COM O DA PLANILHA
        if autor_raw.strip().lower() != autor_sistema_texto.lower():
            print(f"[ALERTA LINHA {linha}] Divergência detectada! Planilha: '{autor_raw}' vs Sistema: '{autor_sistema_texto}'")
            
            # Salva os dados da obra na lista de não indexadas
            obras_nao_indexadas.append({
                "linha": mountain_range,
                "linha": linha,
                "titulo": titulo,
                "subtitulo": subtitulo,
                "autor_planilha": autor_raw,
                "autor_sistema": autor_sistema_texto
            })
            
            try:
                print(f"[LINHA {linha}] Clicando em Cancelar...")
                btn_cancelar = driver.find_element(By.XPATH, "//a[contains(@onclick, 'CatalogingInput.cancelEdit')]")
                driver.execute_script("arguments[0].click();", btn_cancelar)
                time.sleep(1.0)
                
                print(f"[LINHA {linha}] Confirmando cancelamento clicando em 'Sim'...")
                btn_sim = driver.find_element(By.XPATH, "//a[contains(@class, 'button') and text()='Sim']")
                driver.execute_script("arguments[0].click();", btn_sim)
                time.sleep(1.5)
            except Exception as e_cancelar:
                print(f"[AVISO LINHA {linha}] Falha ao tentar fechar caixinha de confirmação: {e_cancelar}")
                
            print(f"[LINHA {linha}] Avançando de linha com segurança.")
            continue
            
        # -------------------------------------------------------------------
        # PASSO 2.3: PREENCHIMENTO DO CAMPO MARC 090 (Número de Chamada)
        # -------------------------------------------------------------------
        if cdd:
            print(f"[LINHA {linha}] Preenchendo Classificação (090 $a) com: {cdd}")
            input_classificacao = driver.find_element(By.XPATH, "//fieldset[@data='090']//div[@data='090' and contains(@class, 'subfield')]//input[@name='a']")
            driver.execute_script("arguments[0].scrollIntoView({block:'center'});", input_classificacao)
            input_classificacao.clear()
            input_classificacao.send_keys(str(cdd))
            
        if codigo_autor:
            print(f"[LINHA {linha}] Preenchendo Código do Autor (090 $b) com: {codigo_autor}")
            input_codigo_autor = driver.find_element(By.XPATH, "//fieldset[@data='090']//div[@data='090' and contains(@class, 'subfield')]//input[@name='b']")
            input_codigo_autor.clear()
            input_codigo_autor.send_keys(str(codigo_autor))
            
        driver.save_screenshot(f"linha_{linha}_02_preenchida.png")
        time.sleep(0.5)
        
        # -------------------------------------------------------------------
        # PASSO 2.4: SALVAMENTO DO REGISTRO
        # -------------------------------------------------------------------
        print(f"[LINHA {linha}] Clicando em Salvar...")
        botao_salvar = driver.find_element(By.XPATH, "//a[contains(@onclick,'CatalogingInput.saveRecord')]")
        driver.execute_script("arguments[0].scrollIntoView(true);", botao_salvar)
        time.sleep(0.5)
        driver.execute_script("arguments[0].click();", botao_salvar)
        
        time.sleep(3.5)
        driver.save_screenshot(f"linha_{linha}_03_salva.png")
        print(f"[LINHA {linha}] Registro atualizado e salvo com sucesso!")

    except Exception as erro:
        print(f"[ERRO LINHA {linha}] Ocorreu uma falha no processamento: {erro}")
        driver.save_screenshot(f"linha_{linha}_erro.png")
        continue

print("\n" + "=" * 60)
print("[INFO] Processo de processamento finalizado!")
driver.quit()

# ---------------------------------------------------------------------------
# PASSO 3: EXIBIÇÃO DO RELATÓRIO DE OBRAS NÃO INDEXADAS
# ---------------------------------------------------------------------------
print("\n" + "X" * 60)
print(f"   RELATÓRIO DE OBRAS NÃO INDEXADAS (TOTAL: {len(obras_nao_indexadas)})")
print("X" * 60)

if obras_nao_indexadas:
    for obra in obras_nao_indexadas:
        print(f"\n-> Linha Planilha: {obra['linha']}")
        print(f"   Título: {obra['titulo']}")
        if obra['subtitulo']:
            print(f"   Subtítulo: {obra['subtitulo']}")
        print(f"   Autor na Planilha: {obra['autor_planilha']}")
        print(f"   Autor no Biblivre: {obra['autor_sistema']}")
else:
    print("\n[SUCESSO] Todas as obras pesquisadas bateram com os autores e foram indexadas!")
print("\n" + "X" * 60)