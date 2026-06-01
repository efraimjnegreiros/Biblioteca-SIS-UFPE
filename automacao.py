# # from selenium import webdriver
# # from selenium.webdriver.common.by import By
# # from selenium.webdriver.common.keys import Keys
# # from openpyxl import load_workbook
# # import time

# # # PLANILHA
# # wb = load_workbook("planilha.xlsx")
# # ws = wb["Sheet1"]

# # # CHROME
# # options = webdriver.ChromeOptions()
# # driver = webdriver.Chrome(options=options)

# # # Ajuste da janela para garantir que todos os elementos fiquem visíveis na tela
# # driver.maximize_window()

# # print("[INFO] Abrindo Biblivre...")
# # driver.get("http://localhost/Biblivre5/")

# # time.sleep(3)

# # # ---------------------------------------------------------------------------
# # # PASSO 1: ROTINA DE LOGIN
# # # ---------------------------------------------------------------------------
# # try:
# #     print("[INFO] Realizando login...")
    
# #     # Preenche o Usuário
# #     campo_usuario = driver.find_element(By.NAME, "username")
# #     campo_usuario.clear()
# #     campo_usuario.send_keys("admin")
    
# #     # Preenche a Senha
# #     campo_senha = driver.find_element(By.NAME, "password")
# #     campo_senha.clear()
# #     campo_senha.send_keys("Bibsis2025#")
    
# #     # Clica no botão Entrar
# #     botao_entrar = driver.find_element(By.XPATH, "//button[contains(@onclick, \"submitForm('login'\")]")
# #     botao_entrar.click()
    
# #     time.sleep(3)
# #     print("[INFO] Login realizado com sucesso!")
# # except Exception as e:
# #     print(f"[ERRO LOGIN] Não foi possível logar: {e}")
# #     driver.save_screenshot("erro_login.png")
# #     driver.quit()
# #     exit()

# # # ---------------------------------------------------------------------------
# # # PASSO 2: NAVEGAÇÃO NO MENU (Catalogação -> Bibliográfica)
# # # ---------------------------------------------------------------------------
# # try:
# #     print("[INFO] Navegando até o menu Catalogação Bibliográfica...")
    
# #     # Clica no menu principal "Catalogação"
# #     menu_catalogacao = driver.find_element(By.CSS_SELECTOR, "li.menu_cataloging")
# #     menu_catalogacao.click()
# #     time.sleep(1)
    
# #     # Clica no submenu "Bibliográfica"
# #     submenu_bibliografica = driver.find_element(By.XPATH, "//li[@data-action='cataloging_bibliographic']/a")
# #     submenu_bibliografica.click()
    
# #     time.sleep(3)
# #     print("[INFO] Página de Catalogação Bibliográfica carregada.")
# # except Exception as e:
# #     print(f"[ERRO NAVEGAÇÃO] Falha ao navegar nos menus: {e}")
# #     driver.save_screenshot("erro_navegacao.png")
# #     driver.quit()
# #     exit()

# # # ---------------------------------------------------------------------------
# # # PASSO 3: LAÇO DE REPETIÇÃO DA PLANILHA
# # # ---------------------------------------------------------------------------
# # for linha in range(3, 5):

# #     print("\n" + "=" * 60)
# #     print(f"[LINHA {linha}] Iniciando processamento")

# #     titulo = ws[f'B{linha}'].value or ""
# #     subtitulo = ws[f'C{linha}'].value or ""
# #     autor = ws[f'D{linha}'].value or ""
# #     editora = ws[f'E{linha}'].value or ""
# #     cidade = ws[f'F{linha}'].value or ""
# #     ano = ws[f'G{linha}'].value or ""
# #     edicao = ws[f'H{linha}'].value or ""  # Nova coluna mapeada com base no seu HTML
# #     assuntos = ws[f'I{linha}'].value or ""
# #     cdd = ws[f'K{linha}'].value or ""

# #     print(f"[LINHA {linha}] Título: {titulo}")

# #     # Tratamento da lista de assuntos separados por ponto e vírgula
# #     lista_assuntos = [
# #         x.strip()
# #         for x in str(assuntos).split(";")
# #         if x.strip()
# #     ]

# #     try:
# #         print(f"[LINHA {linha}] Clicando em Novo Registro")
# #         driver.find_element(By.ID, "new_record_button").click()
# #         time.sleep(2.5)

# #         driver.save_screenshot(f"linha_{linha}_01_novo_registro.png")

# #         # -------------------------------------------------------------------
# #         # PREENCHIMENTO PRECISO CONFORME MAPA DE ABAS E CAMPOS MARC
# #         # -------------------------------------------------------------------

# #         # 1. CDD - Número de Classificação (Coluna K)
# #         if cdd:
# #             print(f"[LINHA {linha}] Preenchendo CDD")
# #             try:
# #                 campo_cdd = driver.find_element(By.XPATH, "//div[@data='082']//input[@name='a']")
# #             except:
# #                 # Fallback caso a div pai não use explicitamente data=082, pega o primeiro input 'a' isolado
# #                 campo_cdd = driver.find_element(By.XPATH, "(//input[@name='a' and contains(@class, 'finput')])[1]")
# #             campo_cdd.clear()
# #             campo_cdd.send_keys(str(cdd))

# #         # 2. AUTOR - Sobrenome e/ou prenome (Coluna D -> MARC 100 $a)
# #         if autor:
# #             print(f"[LINHA {linha}] Preenchendo Autor")
# #             campo_autor = driver.find_element(By.XPATH, "//div[@data='100']//input[@name='a']")
# #             campo_autor.clear()
# #             campo_autor.send_keys(str(autor))
# #             time.sleep(0.5)
# #             campo_autor.send_keys(Keys.ESCAPE)  # Fecha sugestão flutuante

# #         # 3. TÍTULO - Título Principal (Coluna B -> MARC 245 $a)
# #         if titulo:
# #             print(f"[LINHA {linha}] Preenchendo Título")
# #             campo_titulo = driver.find_element(By.XPATH, "//div[@data='245']//input[@name='a']")
# #             campo_titulo.clear()
# #             campo_titulo.send_keys(str(titulo))

# #         # 4. SUBTÍTULO - Títulos paralelos/subtítulos (Coluna C -> MARC 245 $b)
# #         if subtitulo:
# #             print(f"[LINHA {linha}] Preenchendo Subtítulo")
# #             campo_subtitulo = driver.find_element(By.XPATH, "//div[@data='245']//input[@name='b']")
# #             campo_subtitulo.clear()
# #             campo_subtitulo.send_keys(str(subtitulo))

# #         # 5. EDIÇÃO - Indicação da Edição (Coluna H -> MARC 250 $a)
# #         if edicao:
# #             print(f"[LINHA {linha}] Preenchendo Edição")
# #             campo_edicao = driver.find_element(By.XPATH, "//div[@data='250']//input[@name='a']")
# #             campo_edicao.clear()
# #             campo_edicao.send_keys(str(edicao))

# #         # 6. LOCAL - Local de publicação (Coluna F -> MARC 260 $a)
# #         if cidade:
# #             print(f"[LINHA {linha}] Preenchendo Cidade")
# #             campo_cidade = driver.find_element(By.XPATH, "//div[@data='260']//input[@name='a']")
# #             campo_cidade.clear()
# #             campo_cidade.send_keys(str(cidade))

# #         # 7. EDITORA - Nome do editor (Coluna E -> MARC 260 $b)
# #         if editora:
# #             print(f"[LINHA {linha}] Preenchendo Editora")
# #             campo_editora = driver.find_element(By.XPATH, "//div[@data='260']//input[@name='b']")
# #             campo_editora.clear()
# #             campo_editora.send_keys(str(editora))

# #         # 8. ANO - Data de publicação (Coluna G -> MARC 260 $c)
# #         if ano:
# #             print(f"[LINHA {linha}] Preenchendo Ano")
# #             campo_ano = driver.find_element(By.XPATH, "//div[@data='260']//input[@name='c']")
# #             campo_ano.clear()
# #             campo_ano.send_keys(str(ano))

# #         # 9. ASSUNTOS - Vocabulário Controlado (Coluna I -> MARC 650 $a)
# #         if lista_assuntos:
# #             print(f"[LINHA {linha}] Preenchendo Assuntos")
# #             # Encontra as caixas de texto com o autocomplete de vocabulário que estejam visíveis
# #             todos_campos_assunto = driver.find_elements(By.XPATH, "//input[@data-ac='vocabulary' and @name='a']")
# #             campos_assunto_visiveis = [el for el in todos_campos_assunto if el.is_displayed()]
            
# #             for i, assunto_texto in enumerate(lista_assuntos):
# #                 if i >= len(campos_assunto_visiveis):
# #                     print(f"[AVISO LINHA {linha}] Limite de inputs de assunto na tela atingido. Ignorando restantes.")
# #                     break
# #                 campos_assunto_visiveis[i].clear()
# #                 campos_assunto_visiveis[i].send_keys(str(assunto_texto))
# #                 time.sleep(0.5)
# #                 campos_assunto_visiveis[i].send_keys(Keys.ESCAPE)  # Fecha o painel de termos sugeridos

# #         driver.save_screenshot(f"linha_{linha}_02_preenchida.png")
# #         time.sleep(1)

# #         # --- SALVAMENTO SEGURO ---
# #         print(f"[LINHA {linha}] Clicando em Salvar")
        
# #         botao_salvar = driver.find_element(
# #             By.XPATH, 
# #             "//a[contains(@onclick,'CatalogingInput.saveRecord')]"
# #         )
        
# #         # Garante foco visual no botão antes de clicar
# #         driver.execute_script("arguments[0].scrollIntoView(true);", botao_salvar)
# #         time.sleep(0.5)
        
# #         # Dispara o clique via JS para burlar qualquer overlay flutuante de autocompletes abertos
# #         driver.execute_script("arguments[0].click();", botao_salvar)

# #         time.sleep(3.5)  # Tempo de gravação no banco local do Biblivre

# #         driver.save_screenshot(f"linha_{linha}_03_salva.png")
# #         print(f"[LINHA {linha}] Registro salvo com sucesso!")

# #     except Exception as erro:
# #         print(f"[ERRO LINHA {linha}] {erro}")
# #         driver.save_screenshot(f"linha_{linha}_erro.png")
# #         break

# # print("\n[INFO] Processo finalizado")
# # driver.quit()
# from selenium import webdriver
# from selenium.webdriver.common.by import By
# from selenium.webdriver.common.keys import Keys
# from openpyxl import load_workbook
# import time

# # PLANILHA
# wb = load_workbook("planilha.xlsx")
# ws = wb["Sheet1"]

# # CHROME
# options = webdriver.ChromeOptions()
# driver = webdriver.Chrome(options=options)

# # Ajuste da janela para garantir que todos os elementos fiquem visíveis na tela
# driver.maximize_window()

# print("[INFO] Abrindo Biblivre...")
# driver.get("http://localhost/Biblivre5/")

# time.sleep(3)

# # ---------------------------------------------------------------------------
# # PASSO 1: ROTINA DE LOGIN
# # ---------------------------------------------------------------------------
# try:
#     print("[INFO] Realizando login...")
    
#     # Preenche o Usuário
#     campo_usuario = driver.find_element(By.NAME, "username")
#     campo_usuario.clear()
#     campo_usuario.send_keys("admin")
    
#     # Preenche a Senha
#     campo_senha = driver.find_element(By.NAME, "password")
#     campo_senha.clear()
#     campo_senha.send_keys("Bibsis2025#")
    
#     # Clica no botão Entrar
#     botao_entrar = driver.find_element(By.XPATH, "//button[contains(@onclick, \"submitForm('login'\")]")
#     botao_entrar.click()
    
#     time.sleep(3)
#     print("[INFO] Login realizado com sucesso!")
# except Exception as e:
#     print(f"[ERRO LOGIN] Não foi possível logar: {e}")
#     driver.save_screenshot("erro_login.png")
#     driver.quit()
#     exit()

# # ---------------------------------------------------------------------------
# # PASSO 2: NAVEGAÇÃO NO MENU (Catalogação -> Bibliográfica)
# # ---------------------------------------------------------------------------
# try:
#     print("[INFO] Navegando até o menu Catalogação Bibliográfica...")
    
#     # Clica no menu principal "Catalogação"
#     menu_catalogacao = driver.find_element(By.CSS_SELECTOR, "li.menu_cataloging")
#     menu_catalogacao.click()
#     time.sleep(1)
    
#     # Clica no submenu "Bibliográfica"
#     submenu_bibliografica = driver.find_element(By.XPATH, "//li[@data-action='cataloging_bibliographic']/a")
#     submenu_bibliografica.click()
    
#     time.sleep(3)
#     print("[INFO] Página de Catalogação Bibliográfica carregada.")
# except Exception as e:
#     print(f"[ERRO NAVEGAÇÃO] Falha ao navegar nos menus: {e}")
#     driver.save_screenshot("erro_navegacao.png")
#     driver.quit()
#     exit()

# # ---------------------------------------------------------------------------
# # PASSO 3: LAÇO DE REPETIÇÃO DA PLANILHA
# # ---------------------------------------------------------------------------
# for linha in range(3, 5):

#     print("\n" + "=" * 60)
#     print(f"[LINHA {linha}] Iniciando processamento")

#     titulo = ws[f'B{linha}'].value or ""
#     subtitulo = ws[f'C{linha}'].value or ""
#     autor = ws[f'D{linha}'].value or ""
#     editora = ws[f'E{linha}'].value or ""
#     cidade = ws[f'F{linha}'].value or ""
#     ano = ws[f'G{linha}'].value or ""
#     edicao = ws[f'H{linha}'].value or ""
#     assuntos = ws[f'I{linha}'].value or ""
#     cdd = ws[f'K{linha}'].value or ""

#     print(f"[LINHA {linha}] Título: {titulo}")

#     # Tratamento da lista de assuntos separados por ponto e vírgula
#     lista_assuntos = [
#         x.strip()
#         for x in str(assuntos).split(";")
#         if x.strip()
#     ]

#     try:
#         print(f"[LINHA {linha}] Clicando em Novo Registro")
#         driver.find_element(By.ID, "new_record_button").click()
#         time.sleep(2.5)

#         driver.save_screenshot(f"linha_{linha}_01_novo_registro.png")

#         # -------------------------------------------------------------------
#         # PREENCHIMENTO PRECISO CONFORME MAPA DE ABAS E CAMPOS MARC
#         # -------------------------------------------------------------------

#         # 1. CDD - Número de Classificação (Coluna K)
#         if cdd:
#             print(f"[LINHA {linha}] Preenchendo CDD")
#             try:
#                 campo_cdd = driver.find_element(By.開Xpath, "//div[@data='082']//input[@name='a']")
#             except:
#                 campo_cdd = driver.find_element(By.XPATH, "(//input[@name='a' and contains(@class, 'finput')])[1]")
#             campo_cdd.clear()
#             campo_cdd.send_keys(str(cdd))

#         # 2. AUTOR - Sobrenome e/ou prenome (Coluna D -> MARC 100 $a)
#         if autor:
#             print(f"[LINHA {linha}] Preenchendo Autor")
#             campo_autor = driver.find_element(By.XPATH, "//div[@data='100']//input[@name='a']")
#             campo_autor.clear()
#             campo_autor.send_keys(str(autor))
#             time.sleep(0.5)
#             campo_autor.send_keys(Keys.ESCAPE)

#         # 3. TÍTULO - Título Principal (Coluna B -> MARC 245 $a)
#         if titulo:
#             print(f"[LINHA {linha}] Preenchendo Título")
#             campo_titulo = driver.find_element(By.XPATH, "//div[@data='245']//input[@name='a']")
#             campo_titulo.clear()
#             campo_titulo.send_keys(str(titulo))

#         # 4. SUBTÍTULO - Títulos paralelos/subtítulos (Coluna C -> MARC 245 $b)
#         if subtitulo:
#             print(f"[LINHA {linha}] Preenchendo Subtítulo")
#             campo_subtitulo = driver.find_element(By.XPATH, "//div[@data='245']//input[@name='b']")
#             campo_subtitulo.clear()
#             campo_subtitulo.send_keys(str(subtitulo))

#         # 5. EDIÇÃO - Indicação da Edição (Coluna H -> MARC 250 $a)
#         if edicao:
#             print(f"[LINHA {linha}] Preenchendo Edição")
#             campo_edicao = driver.find_element(By.XPATH, "//div[@data='250']//input[@name='a']")
#             campo_edicao.clear()
#             campo_edicao.send_keys(str(edicao))

#         # 6. LOCAL - Local de publicação (Coluna F -> MARC 260 $a)
#         if cidade:
#             print(f"[LINHA {linha}] Preenchendo Cidade")
#             campo_cidade = driver.find_element(By.XPATH, "//div[@data='260']//input[@name='a']")
#             campo_cidade.clear()
#             campo_cidade.send_keys(str(cidade))

#         # 7. EDITORA - Nome do editor (Coluna E -> MARC 260 $b)
#         if editora:
#             print(f"[LINHA {linha}] Preenchendo Editora")
#             campo_editora = driver.find_element(By.XPATH, "//div[@data='260']//input[@name='b']")
#             campo_editora.clear()
#             campo_editora.send_keys(str(editora))

#         # 8. ANO - Data de publicação (Coluna G -> MARC 260 $c)
#         if ano:
#             print(f"[LINHA {linha}] Preenchendo Ano")
#             campo_ano = driver.find_element(By.XPATH, "//div[@data='260']//input[@name='c']")
#             campo_ano.clear()
#             campo_ano.send_keys(str(ano))

#         # 9. ASSUNTOS - Múltiplos subcampos sequenciais (Coluna I -> MARC 650 de 'a' até 'j')
#         if lista_assuntos:
#             print(f"[LINHA {linha}] Preenchendo Assuntos...")
#             # Modificado: captura todos os inputs de vocabulário (independentemente do name="a", "b", etc.)
#             todos_campos_assunto = driver.find_elements(By.XPATH, "//input[@data-ac='vocabulary']")
#             campos_assunto_visiveis = [el for el in todos_campos_assunto if el.is_displayed()]
            
#             for i, assunto_texto in enumerate(lista_assuntos):
#                 if i >= len(campos_assunto_visiveis):
#                     print(f"[AVISO LINHA {linha}] Limite de inputs de assunto na tela atingido ({len(campos_assunto_visiveis)}). Ignorando restantes.")
#                     break
                
#                 print(f"  -> Preenchendo assunto {i+1}: {assunto_texto}")
#                 campos_assunto_visiveis[i].clear()
#                 campos_assunto_visiveis[i].send_keys(str(assunto_texto))
#                 time.sleep(0.5)
#                 campos_assunto_visiveis[i].send_keys(Keys.ESCAPE)  # Fecha o popup de sugestões do Biblivre

#         driver.save_screenshot(f"linha_{linha}_02_preenchida.png")
#         time.sleep(1)

#         # --- SALVAMENTO SEGURO ---
#         print(f"[LINHA {linha}] Clicando em Salvar")
        
#         botao_salvar = driver.find_element(
#             By.XPATH, 
#             "//a[contains(@onclick,'CatalogingInput.saveRecord')]"
#         )
        
#         driver.execute_script("arguments[0].scrollIntoView(true);", botao_salvar)
#         time.sleep(0.5)
#         driver.execute_script("arguments[0].click();", botao_salvar)

#         time.sleep(3.5)

#         driver.save_screenshot(f"linha_{linha}_03_salva.png")
#         print(f"[LINHA {linha}] Registro salvo com sucesso!")

#     except Exception as erro:
#         print(f"[ERRO LINHA {linha}] {erro}")
#         driver.save_screenshot(f"linha_{linha}_erro.png")
#         break

# print("\n[INFO] Processo finalizado")
# driver.quit()
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from openpyxl import load_workbook
import time

# PLANILHA
wb = load_workbook("planilha.xlsx")
ws = wb["Sheet1"]

# CHROME
options = webdriver.ChromeOptions()
driver = webdriver.Chrome(options=options)

# Ajuste da janela para garantir que todos os elementos fiquem visíveis na tela
driver.maximize_window()

print("[INFO] Abrindo Biblivre...")
driver.get("http://localhost/Biblivre5/")

time.sleep(3)

# ---------------------------------------------------------------------------
# PASSO 1: ROTINA DE LOGIN
# ---------------------------------------------------------------------------
try:
    print("[INFO] Realizando login...")
    
    # Preenche o Usuário
    campo_usuario = driver.find_element(By.NAME, "username")
    campo_usuario.clear()
    campo_usuario.send_keys("admin")
    
    # Preenche a Senha
    campo_senha = driver.find_element(By.NAME, "password")
    campo_senha.clear()
    campo_senha.send_keys("Bibsis2025#")
    
    # Clica no botão Entrar
    botao_entrar = driver.find_element(By.XPATH, "//button[contains(@onclick, \"submitForm('login'\")]")
    botao_entrar.click()
    
    time.sleep(3)
    print("[INFO] Login realizado com sucesso!")
except Exception as e:
    print(f"[ERRO LOGIN] Não foi possível logar: {e}")
    driver.save_screenshot("erro_login.png")
    driver.quit()
    exit()

# ---------------------------------------------------------------------------
# PASSO 2: NAVEGAÇÃO NO MENU (Catalogação -> Bibliográfica)
# ---------------------------------------------------------------------------
try:
    print("[INFO] Navegando até o menu Catalogação Bibliográfica...")
    
    # Clica no menu principal "Catalogação"
    menu_catalogacao = driver.find_element(By.CSS_SELECTOR, "li.menu_cataloging")
    menu_catalogacao.click()
    time.sleep(1)
    
    # Clica no submenu "Bibliográfica"
    submenu_bibliografica = driver.find_element(By.XPATH, "//li[@data-action='cataloging_bibliographic']/a")
    submenu_bibliografica.click()
    
    time.sleep(3)
    print("[INFO] Página de Catalogação Bibliográfica carregada.")
except Exception as e:
    print(f"[ERRO NAVEGAÇÃO] Falha ao navegar nos menus: {e}")
    driver.save_screenshot("erro_navegacao.png")
    driver.quit()
    exit()

# ---------------------------------------------------------------------------
# PASSO 3: LAÇO DE REPETIÇÃO DA PLANILHA
# ---------------------------------------------------------------------------
for linha in range(500, 601):

    print("\n" + "=" * 60)
    print(f"[LINHA {linha}] Iniciando processamento")

    titulo = ws[f'B{linha}'].value or ""
    subtitulo = ws[f'C{linha}'].value or ""
    autor = ws[f'D{linha}'].value or ""
    editora = ws[f'E{linha}'].value or ""
    cidade = ws[f'F{linha}'].value or ""
    ano = ws[f'G{linha}'].value or ""
    edicao = ws[f'H{linha}'].value or ""
    assuntos = ws[f'I{linha}'].value or ""
    cdd = ws[f'K{linha}'].value or ""

    print(f"[LINHA {linha}] Título: {titulo}")

    try:
        print(f"[LINHA {linha}] Clicando em Novo Registro")
        driver.find_element(By.ID, "new_record_button").click()
        time.sleep(2.5)

        driver.save_screenshot(f"linha_{linha}_01_novo_registro.png")

        # -------------------------------------------------------------------
        # PREENCHIMENTO PRECISO CONFORME MAPA DE ABAS E CAMPOS MARC
        # -------------------------------------------------------------------

        # 1. CDD - Número de Classificação (Coluna K)
        if cdd:
            print(f"[LINHA {linha}] Preenchendo CDD")
            try:
                campo_cdd = driver.find_element(By.XPATH, "//div[@data='082']//input[@name='a']")
            except:
                campo_cdd = driver.find_element(By.XPATH, "(//input[@name='a' and contains(@class, 'finput')])[1]")
            campo_cdd.clear()
            campo_cdd.send_keys(str(cdd))

        # 2. AUTOR - Sobrenome e/ou prenome (Coluna D -> MARC 100 $a)
        if autor:
            print(f"[LINHA {linha}] Preenchendo Autor")
            campo_autor = driver.find_element(By.XPATH, "//div[@data='100']//input[@name='a']")
            campo_autor.clear()
            campo_autor.send_keys(str(autor))
            time.sleep(0.5)
            campo_autor.send_keys(Keys.ESCAPE)

        # 3. TÍTULO - Título Principal (Coluna B -> MARC 245 $a)
        if titulo:
            print(f"[LINHA {linha}] Preenchendo Título")
            campo_titulo = driver.find_element(By.XPATH, "//div[@data='245']//input[@name='a']")
            campo_titulo.clear()
            campo_titulo.send_keys(str(titulo))

        # 4. SUBTÍTULO - Títulos paralelos/subtítulos (Coluna C -> MARC 245 $b)
        if subtitulo:
            print(f"[LINHA {linha}] Preenchendo Subtítulo")
            campo_subtitulo = driver.find_element(By.XPATH, "//div[@data='245']//input[@name='b']")
            campo_subtitulo.clear()
            campo_subtitulo.send_keys(str(subtitulo))

        # 5. EDIÇÃO - Indicação da Edição (Coluna H -> MARC 250 $a)
        if edicao:
            print(f"[LINHA {linha}] Preenchendo Edição")
            campo_edicao = driver.find_element(By.XPATH, "//div[@data='250']//input[@name='a']")
            campo_edicao.clear()
            campo_edicao.send_keys(str(edicao))

        # 6. LOCAL - Local de publicação (Coluna F -> MARC 260 $a)
        if cidade:
            print(f"[LINHA {linha}] Preenchendo Cidade")
            campo_cidade = driver.find_element(By.XPATH, "//div[@data='260']//input[@name='a']")
            campo_cidade.clear()
            campo_cidade.send_keys(str(cidade))

        # 7. EDITORA - Nome do editor (Coluna E -> MARC 260 $b)
        if editora:
            print(f"[LINHA {linha}] Preenchendo Editora")
            campo_editora = driver.find_element(By.XPATH, "//div[@data='260']//input[@name='b']")
            campo_editora.clear()
            campo_editora.send_keys(str(editora))

        # 8. ANO - Data de publicação (Coluna G -> MARC 260 $c)
        if ano:
            print(f"[LINHA {linha}] Preenchendo Ano")
            campo_ano = driver.find_element(By.XPATH, "//div[@data='260']//input[@name='c']")
            campo_ano.clear()
            campo_ano.send_keys(str(ano))

        # 9. ASSUNTOS (650$a até 650$j)
        if assuntos:
            lista_assuntos = [
                assunto.strip()
                for assunto in str(assuntos).split(";")
                if assunto.strip()
            ]

            print(f"[LINHA {linha}] {len(lista_assuntos)} assuntos encontrados")

            letras = ["a", "b", "c", "d", "e", "f", "g", "h", "i", "j"]

            for indice, assunto_texto in enumerate(lista_assuntos):
                if indice >= len(letras):
                    print("[AVISO] Máximo de 10 assuntos atingido")
                    break

                try:
                    letra = letras[indice]
                    campo = driver.find_element(
                        By.XPATH,
                        f"//input[@data-ac='vocabulary' and @name='{letra}']"
                    )

                    driver.execute_script(
                        "arguments[0].scrollIntoView({block:'center'});",
                        campo
                    )

                    campo.clear()
                    campo.send_keys(assunto_texto)

                    print(f"  -> 650${letra}: {assunto_texto}")
                    time.sleep(0.3)
                    campo.send_keys(Keys.ESCAPE)

                except Exception as erro_assunto:
                    print(f"[ERRO ASSUNTO {indice+1}] {erro_assunto}")

        driver.save_screenshot(f"linha_{linha}_02_preenchida.png")
        time.sleep(1)

        # --- SALVAMENTO SEGURO ---
        print(f"[LINHA {linha}] Clicando em Salvar")
        
        botao_salvar = driver.find_element(
            By.XPATH, 
            "//a[contains(@onclick,'CatalogingInput.saveRecord')]"
        )
        
        driver.execute_script("arguments[0].scrollIntoView(true);", botao_salvar)
        time.sleep(0.5)
        driver.execute_script("arguments[0].click();", botao_salvar)

        time.sleep(3.5)

        driver.save_screenshot(f"linha_{linha}_03_salva.png")
        print(f"[LINHA {linha}] Registro salvo com sucesso!")

    except Exception as erro:
        print(f"[ERRO LINHA {linha}] {erro}")
        driver.save_screenshot(f"linha_{linha}_erro.png")
        break

print("\n[INFO] Processo finalizado")
driver.quit()